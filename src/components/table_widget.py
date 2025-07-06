#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
现代化表格组件模块
负责处理发票表格的显示和交互
"""

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from PySide6.QtWidgets import (QWidget, QGroupBox, QVBoxLayout, QTableView,
                               QHeaderView, QMessageBox, QLabel,
                               QHBoxLayout, QPushButton, QFileDialog, QStyledItemDelegate, QCheckBox, QMenu)
from PySide6.QtCore import Qt, Signal, QAbstractTableModel, QModelIndex, QSortFilterProxyModel
from PySide6.QtGui import QColor, QFont, QAction

class InvoiceTableModel(QAbstractTableModel):
    """发票数据的自定义表格模型"""
    check_state_changed = Signal()

    def __init__(self, data=None, headers=None, parent=None):
        super().__init__(parent)
        self._data = data or []
        self._headers = headers or []
        self._check_states = [Qt.CheckState.Unchecked] * self.rowCount()

    def rowCount(self, parent=QModelIndex()):
        return len(self._data)

    def columnCount(self, parent=QModelIndex()):
        return len(self._headers)

    def flags(self, index):
        """设置单元格的标志，使第一列可选中山"""
        if not index.isValid():
            return super().flags(index)
        
        # 允许第一列是可选中的
        if index.column() == 0:
            return super().flags(index) | Qt.ItemFlag.ItemIsUserCheckable
            
        return super().flags(index)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        
        row, col = index.row(), index.column()

        if role == Qt.ItemDataRole.CheckStateRole and col == 0:
            return self._check_states[row]
        
        if role == Qt.ItemDataRole.DisplayRole:
            value = self._data[row][col]
            if col == 5: # 金额
                try:
                    return f"¥{float(str(value).replace('¥','').replace(',','')):.2f}"
                except (ValueError, TypeError):
                    return "¥0.00"
            return str(value)
        
        if role == Qt.ItemDataRole.TextAlignmentRole:
            return Qt.AlignmentFlag.AlignCenter

        if role == Qt.ItemDataRole.UserRole: # For sorting/filtering raw data
            return self._data[row][col]
            
        return None

    def setData(self, index, value, role=Qt.ItemDataRole.EditRole):
        if not index.isValid() or role != Qt.ItemDataRole.CheckStateRole or index.column() != 0:
            return super().setData(index, value, role)

        self._check_states[index.row()] = Qt.CheckState(value)
        self.dataChanged.emit(index, index, [role])
        self.check_state_changed.emit()
        return True

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole and orientation == Qt.Orientation.Horizontal:
            return self._headers[section]
        return None
        
    def set_data(self, data):
        self.beginResetModel()
        self._data = data
        self._check_states = [Qt.CheckState.Unchecked] * self.rowCount()
        self.endResetModel()

    def get_row_data(self, row):
        if 0 <= row < len(self._data):
            return self._data[row]
        return None

    def get_checked_row_indices(self):
        """获取所有选中行的索引"""
        return [i for i, state in enumerate(self._check_states) if state == Qt.CheckState.Checked]

    def get_checked_row_data(self):
        """获取所有选中行的数据"""
        checked_indices = self.get_checked_row_indices()
        return [self.get_row_data(i) for i in checked_indices]

    def set_all_checked(self, is_checked):
        """设置所有行的选中状态"""
        if not self._data:
            return
        
        state = Qt.CheckState.Checked if is_checked else Qt.CheckState.Unchecked
        self._check_states = [state] * self.rowCount()
        
        top_left = self.index(0, 0)
        bottom_right = self.index(self.rowCount() - 1, 0)
        self.dataChanged.emit(top_left, bottom_right, [Qt.ItemDataRole.CheckStateRole])
        self.check_state_changed.emit()

class StatusDelegate(QStyledItemDelegate):
    """用于'状态'列的自定义委托"""
    def paint(self, painter, option, index):
        status = index.data()
        if not status:
            super().paint(painter, option, index)
            return

        color_map = {
            "已使用": ("#dcfce7", "#166534"), # Green
            "未使用": ("#f3f4f6", "#374151"), # Gray
        }
        
        painter.save()
        
        if status in color_map:
            bg_color, fg_color = color_map[status]
            
            # Draw background
            painter.fillRect(option.rect, QColor(bg_color))
            
            # Draw text
            painter.setPen(QColor(fg_color))
            font = QFont(option.font)
            font.setBold(True)
            painter.setFont(font)
            
            text_rect = painter.fontMetrics().boundingRect(option.rect, Qt.AlignmentFlag.AlignCenter, status)
            painter.drawText(text_rect, status)
            
        else:
            super().paint(painter, option, index)
            
        painter.restore()
        

class InvoiceTableWidget(QWidget):
    """现代化发票表格组件"""
    
    selection_changed = Signal(int, int)
    row_double_clicked = Signal(int)
    data_changed = Signal()
    checked_rows_changed = Signal(list)
    
    # 为上下文菜单操作添加新信号
    edit_requested = Signal(int)  # 源模型中的行索引
    delete_requested = Signal(list)  # 源模型中的行索引
    status_change_requested = Signal(list, str)  # 行索引列表, 新状态
    print_requested = Signal(list) # 源模型中的行索引列表
    
    def __init__(self, parent=None):
        super().__init__(parent)
        self.all_data = [] # 存储当前用户的所有数据
        
        # 定义表头
        self.headers = [
            "ID", "发票代码", "发票号码", "供应商", "开票日期", "金额", 
            "发票类型", "项目名称", "分类", "备注", "状态"
        ]
        
        self.setup_ui()
        self.connect_signals()

    def setup_ui(self):
        """设置用户界面"""
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        
        table_group = QGroupBox("📋 发票列表")
        table_layout = QVBoxLayout(table_group)
        table_layout.setContentsMargins(20, 25, 20, 20)
        
        self.create_table_header(table_layout)
        
        self.table = QTableView()
        
        self.model = InvoiceTableModel(headers=self.headers)
        self.table.setModel(self.model)

        # 隐藏ID列
        self.table.setColumnHidden(0, True)
        self.table.verticalHeader().setVisible(False)

        self.table.setSelectionBehavior(QTableView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QTableView.SelectionMode.ExtendedSelection)
        self.table.setAlternatingRowColors(True)
        self.table.setSortingEnabled(True)
        self.table.setEditTriggers(QTableView.EditTrigger.NoEditTriggers)
        self.table.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAsNeeded)
        
        # 设置上下文菜单策略
        self.table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        
        # 设置状态列的委托
        status_delegate = StatusDelegate(self.table)
        self.table.setItemDelegateForColumn(10, status_delegate)
        
        self.table.setMinimumHeight(400)
        self.setup_column_widths()
        
        table_layout.addWidget(self.table)
        self.create_table_footer(table_layout)
        layout.addWidget(table_group)

    def create_table_header(self, parent_layout):
        """创建表格头部信息栏"""
        header_widget = QWidget()
        header_layout = QHBoxLayout(header_widget)
        header_layout.setContentsMargins(0, 0, 0, 0)
        
        self.select_all_checkbox = QCheckBox("全选")
        header_layout.addWidget(self.select_all_checkbox)

        self.stats_label = QLabel("共 0 条记录")
        header_layout.addWidget(self.stats_label)
        header_layout.addStretch()

        self.export_button = QPushButton("导出选中")
        self.export_button.setEnabled(False)
        header_layout.addWidget(self.export_button)

        parent_layout.addWidget(header_widget)

    def create_table_footer(self, parent_layout):
        """创建表格底部信息栏"""
        footer_widget = QWidget()
        footer_layout = QHBoxLayout(footer_widget)
        footer_layout.setContentsMargins(0, 0, 0, 0)
        
        self.status_stats_label = QLabel("状态统计: ...")
        self.amount_stats_label = QLabel("总金额: ...")
        
        footer_layout.addWidget(self.status_stats_label)
        footer_layout.addStretch()
        footer_layout.addWidget(self.amount_stats_label)
        parent_layout.addWidget(footer_widget)
        
    def setup_column_widths(self):
        """设置列宽"""
        header = self.table.horizontalHeader()
        header.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.table.setColumnWidth(1, 150)  # 发票代码
        self.table.setColumnWidth(2, 200)  # 发票号码
        self.table.setColumnWidth(3, 220)  # 供应商
        self.table.setColumnWidth(4, 120)  # 开票日期
        self.table.setColumnWidth(5, 120)  # 金额
        self.table.setColumnWidth(6, 150)  # 发票类型
        self.table.setColumnWidth(7, 250)  # 项目名称
        self.table.setColumnWidth(8, 120)  # 分类
        self.table.setColumnWidth(9, 200)  # 备注
        self.table.setColumnWidth(10, 80)   # 状态
        
        # 允许最后一列拉伸以填充剩余空间
        header.setStretchLastSection(True)

    def connect_signals(self):
        """连接信号和槽"""
        self.table.selectionModel().selectionChanged.connect(self.on_selection_changed)
        self.table.doubleClicked.connect(self.on_item_double_clicked)
        self.select_all_checkbox.stateChanged.connect(self.on_select_all_toggled)
        self.model.check_state_changed.connect(self.on_check_state_changed)
        self.export_button.clicked.connect(self.on_export_checked)
        self.table.customContextMenuRequested.connect(self.show_context_menu)
    
    def on_selection_changed(self, selected, deselected):
        """处理选择变化"""
        if self.table.selectionModel().hasSelection():
            proxy_index = self.table.selectionModel().selectedRows()[0]
            # 将视图索引映射到源模型索引以处理排序
            model = self.table.model()
            if isinstance(model, QSortFilterProxyModel):
                source_index = model.mapToSource(proxy_index)
                self.selection_changed.emit(source_index.row(), source_index.column())
            else:
                self.selection_changed.emit(proxy_index.row(), proxy_index.column())
        else:
            self.selection_changed.emit(-1, -1)

    def on_item_double_clicked(self, index):
        """处理项目双击"""
        if index.isValid():
            # 将视图索引映射到源模型索引以处理排序
            model = self.table.model()
            if isinstance(model, QSortFilterProxyModel):
                source_index = model.mapToSource(index)
                self.row_double_clicked.emit(source_index.row())
            else:
                self.row_double_clicked.emit(index.row())

    def on_select_all_toggled(self, state):
        """处理"全选"复选框状态变化"""
        is_checked = (state == Qt.CheckState.Checked.value)
        self.model.set_all_checked(is_checked)

    def on_check_state_changed(self):
        """处理行选中状态变化"""
        checked_rows = self.model.get_checked_row_indices()
        self.checked_rows_changed.emit(checked_rows)
        self.export_button.setEnabled(len(checked_rows) > 0)
        
        # 更新"全选"复选框的状态
        self.select_all_checkbox.blockSignals(True)
        if not self.model.rowCount():
            self.select_all_checkbox.setCheckState(Qt.CheckState.Unchecked)
        elif len(checked_rows) == self.model.rowCount():
            self.select_all_checkbox.setCheckState(Qt.CheckState.Checked)
        elif len(checked_rows) > 0:
            self.select_all_checkbox.setCheckState(Qt.CheckState.PartiallyChecked)
        else:
            self.select_all_checkbox.setCheckState(Qt.CheckState.Unchecked)
        self.select_all_checkbox.blockSignals(False)

    def on_export_checked(self):
        """导出所有选中的行"""
        checked_data = self.model.get_checked_row_data()
        
        if not checked_data:
            QMessageBox.information(self, "无选中项", "请先选择要导出的发票。")
            return
            
        # 将列表数据转换为字典列表以供导出函数使用
        data_to_export = []
        for row_list in checked_data:
            if not row_list:  # 添加空值检查
                continue
            row_dict = {}
            for i, header in enumerate(self.headers):
                key = header_to_key(header)
                if key:
                    row_dict[key] = row_list[i]
            data_to_export.append(row_dict)
            
        self.export_to_excel(data_to_export)

    def set_table_data(self, data):
        """加载数据到表格"""
        self.all_data = data
        self.table.setSortingEnabled(False)
        self.model.set_data(data)
        self.table.setSortingEnabled(True)
        self.update_statistics()
        self.data_changed.emit()
        self.on_check_state_changed()

    def get_current_row(self):
        """获取当前选中的行索引"""
        if not self.table.selectionModel().hasSelection():
            return -1
        proxy_index = self.table.selectionModel().selectedRows()[0]
        # 将视图索引映射到源模型索引以处理排序
        model = self.table.model()
        if isinstance(model, QSortFilterProxyModel):
            source_index = model.mapToSource(proxy_index)
            return source_index.row()
        return proxy_index.row()

    def get_selected_rows(self):
        """获取当前选中的所有行索引 (源模型)"""
        if not self.table.selectionModel().hasSelection():
            return []
        
        selected_proxy_indices = self.table.selectionModel().selectedRows()
        
        model = self.table.model()
        if isinstance(model, QSortFilterProxyModel):
            source_rows = [model.mapToSource(proxy_index).row() for proxy_index in selected_proxy_indices]
        else:
            source_rows = [proxy_index.row() for proxy_index in selected_proxy_indices]
            
        return source_rows

    def get_row_data(self, row):
        """获取指定行的数据"""
        return self.model.get_row_data(row)

    def get_cell_data(self, row, col):
        """获取指定单元格的数据"""
        if 0 <= row < self.model.rowCount() and 0 <= col < self.model.columnCount():
            source_index = self.model.index(row, col)
            return self.model.data(source_index, role=Qt.ItemDataRole.UserRole)
        return None

    def get_row_count(self):
        """获取当前表格的总行数"""
        return self.model.rowCount()
        
    def update_statistics(self):
        """更新统计信息"""
        total_rows = self.get_row_count()
        
        status_counts = {"未使用": 0, "已使用": 0}
        total_amount = 0.0

        for row in range(total_rows):
            status_value = self.model.data(self.model.index(row, 10))
            status = str(status_value) if status_value is not None else ""
            if status in status_counts:
                status_counts[status] += 1
            
            try:
                amount_str = str(self.model.data(self.model.index(row, 5), role=Qt.ItemDataRole.UserRole) or "0")
                amount = float(amount_str.replace("¥", "").replace(",", ""))
                total_amount += amount
            except (ValueError, TypeError):
                pass
        
        self.stats_label.setText(f"共 {total_rows} 条记录")
        self.status_stats_label.setText(
            f"状态: 未使用({status_counts['未使用']}) | "
            f"已使用({status_counts['已使用']})"
        )
        self.amount_stats_label.setText(f"显示总金额: ¥{total_amount:,.2f}")

    def export_to_excel(self, data_to_export):
        """将数据导出到 Excel 文件"""
        if not data_to_export:
            QMessageBox.warning(self, "无数据", "没有可导出的数据。")
            return

        path, _ = QFileDialog.getSaveFileName(self, "保存文件", "", "Excel 文件 (*.xlsx)")
        if not path:
            return

        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            if not sheet:
                QMessageBox.critical(self, "导出失败", "无法创建 Excel 工作表。")
                return

            sheet.title = "发票记录"
            
            export_headers = [h for h in self.headers if h not in ["ID"]]
            sheet.append(export_headers)
            
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
            
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for inv_dict in data_to_export:
                row_data = []
                for header in export_headers:
                    key = header_to_key(header)
                    value = inv_dict.get(key, '')
                    
                    # 确保金额是数字类型
                    if key == 'amount':
                        try:
                            value = float(value)
                        except (ValueError, TypeError):
                            value = 0.0
                    
                    # 确保所有其他值都是字符串或数字
                    if not isinstance(value, (str, int, float, type(None))):
                        value = str(value)
                        
                    row_data.append(value)
                sheet.append(row_data)

            for i, column_cells in enumerate(sheet.columns, 1):
                try:
                    # 增加对列宽计算的健壮性
                    max_len = 0
                    for cell in column_cells:
                        if cell.value:
                            # 将数字转换为字符串来计算长度
                            cell_text = str(cell.value)
                            if len(cell_text) > max_len:
                                max_len = len(cell_text)
                    
                    # 如果列有内容，则设置宽度，否则使用默认值
                    if max_len > 0:
                        sheet.column_dimensions[get_column_letter(i)].width = max_len + 2
                    else:
                        sheet.column_dimensions[get_column_letter(i)].width = 10
                except Exception:
                    # 如果出现任何错误，设置一个默认宽度
                    sheet.column_dimensions[get_column_letter(i)].width = 10

            workbook.save(path)
            QMessageBox.information(self, "导出成功", f"数据已成功导出到:\n{path}")

        except Exception as e:
            QMessageBox.critical(self, "导出失败", f"导出过程中发生错误: {e}")

    def handle_delete_request(self, rows):
        """处理删除请求并显示确认对话框"""
        if not rows:
            return

        count = len(rows)
        if count == 1:
            row_data = self.model.get_row_data(rows[0])
            invoice_number = row_data[2] if row_data else "未知"
            message = f"您确定要删除发票号码为 '{invoice_number}' 的记录吗？\n此操作无法撤销。"
        else:
            message = f"您确定要删除选中的 {count} 条记录吗？\n此操作无法撤销。"

        reply = QMessageBox.question(
            self,
            "确认删除",
            message,
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if reply == QMessageBox.StandardButton.Yes:
            self.delete_requested.emit(rows)

    def show_context_menu(self, position):
        """显示上下文菜单"""
        selected_rows = self.get_selected_rows()
        if not selected_rows:
            return

        num_selected = len(selected_rows)
        menu = QMenu(self)

        # 编辑操作 (仅对单选有效)
        edit_action = QAction(f"✏️ 编辑", self)
        edit_action.triggered.connect(lambda: self.edit_requested.emit(selected_rows[0]))
        edit_action.setEnabled(num_selected == 1)
        menu.addAction(edit_action)

        # 删除操作
        delete_action = QAction(f"❌ 删除 ({num_selected} 项)", self)
        delete_action.triggered.connect(lambda: self.handle_delete_request(selected_rows))
        menu.addAction(delete_action)

        # 打印操作
        print_action = QAction(f"🖨️ 打印 ({num_selected} 项)", self)
        print_action.triggered.connect(lambda: self.print_requested.emit(selected_rows))
        menu.addAction(print_action)

        menu.addSeparator()

        # 更改状态操作 (子菜单)
        status_menu = menu.addMenu(f"🔁 更改状态 ({num_selected} 项)")
        mark_used_action = QAction("🟢 标记为已使用", self)
        mark_used_action.triggered.connect(lambda: self.status_change_requested.emit(selected_rows, "已使用"))
        status_menu.addAction(mark_used_action)
        
        mark_unused_action = QAction("⚪️ 标记为未使用", self)
        mark_unused_action.triggered.connect(lambda: self.status_change_requested.emit(selected_rows, "未使用"))
        status_menu.addAction(mark_unused_action)
        
        # 显示菜单
        menu.exec(self.table.viewport().mapToGlobal(position))

def header_to_key(header):
    """将显示的表头转换为字典的键。"""
    mapping = {
        "发票号码": "invoice_number",
        "供应商": "supplier",
        "开票日期": "issue_date",
        "金额": "amount",
        "发票类型": "invoice_type",
        "项目名称": "item_name",
        "分类": "category",
        "备注": "notes",
        "状态": "status",
    }
    return mapping.get(header, '') 